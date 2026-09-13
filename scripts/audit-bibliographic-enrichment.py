#!/usr/bin/env python3
"""Pilot conservative Crossref enrichment for PhilPapers bibliographic metadata.

This is an audit, not a production data build. It draws a deterministic sample
across publication types and periods, queries Crossref, applies a deliberately
strict title/author/year rule, and records enough evidence for manual review.
"""

from __future__ import annotations

import argparse
import html
import json
import random
import re
import time
import unicodedata
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from difflib import SequenceMatcher
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PAPER_COUNT = 69_400
USER_AGENT = "philpapers-atlas-metadata-audit/0.1 (https://github.com/davidfreeborn/philpapers-atlas)"


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value).casefold())
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def normalize_type(value: object) -> str:
    text = normalized(value)
    if "review" in text:
        return "Review"
    if "book" in text or "книга" in clean_text(value).casefold():
        return "Book"
    if "article" in text or "articolo" in text:
        return "Article"
    return "Other"


def parse_year(value: object) -> int | None:
    match = re.search(r"(?<!\d)(1[5-9]\d{2}|20\d{2})(?!\d)", clean_text(value))
    return int(match.group(1)) if match else None


def period(value: object) -> str:
    text = normalized(value)
    if "manuscript" in text or "forthcoming" in text:
        return "Manuscript/forthcoming"
    year = parse_year(value)
    if year is None:
        return "Undated"
    if year < 1990:
        return "Before 1990"
    if year < 2000:
        return "1990s"
    if year < 2010:
        return "2000s"
    if year < 2020:
        return "2010s"
    return "2020 onward"


def creator_surname(value: object) -> str:
    text = normalized(value)
    if not text:
        return ""
    raw = clean_text(value)
    if "," in raw:
        return normalized(raw.split(",", 1)[0])
    return text.split()[-1]


def load_records(source: Path) -> list[dict]:
    english = np.load(source / "english_indices.npy", allow_pickle=False)
    paths = [
        line.strip()
        for line in (source / "document_ids.txt").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    ids = [Path(paths[int(index)]).stem for index in english]
    if len(ids) != PAPER_COUNT:
        raise ValueError("Unexpected cohort size")
    needed = set(ids)
    workbook = load_workbook(source / "philpapers_metadata.xlsx", read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    columns = {name: index for index, name in enumerate(headers)}
    records: dict[str, dict] = {}
    for values in rows:
        paper_id = clean_text(values[columns["_id"]])
        if paper_id not in needed:
            continue
        record = {
            "id": paper_id,
            "title": clean_text(values[columns["title"]]),
            "creator": clean_text(values[columns["creator"]]),
            "date": clean_text(values[columns["date"]]),
            "type": normalize_type(values[columns["type"]]),
            "description_length": len(clean_text(values[columns["description"]])),
        }
        previous = records.get(paper_id)
        if previous is None or record["description_length"] > previous["description_length"]:
            records[paper_id] = record
    workbook.close()
    return [records[paper_id] for paper_id in ids if paper_id in records]


def draw_sample(records: list[dict], per_stratum: int, seed: int) -> list[dict]:
    randomizer = random.Random(seed)
    strata: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for record in records:
        if not record["title"]:
            continue
        strata[(record["type"], period(record["date"]))].append(record)
    sample = []
    for key in sorted(strata):
        if key[0] == "Other" or key[1] == "Undated":
            continue
        candidates = strata[key]
        sample.extend(randomizer.sample(candidates, min(per_stratum, len(candidates))))
    randomizer.shuffle(sample)
    return sample


def crossref_items(record: dict, rows: int = 5) -> list[dict]:
    query = " ".join(value for value in (record["title"], record["creator"], record["date"]) if value)
    parameters = urlencode(
        {
            "query.bibliographic": query,
            "rows": rows,
            "select": "DOI,title,author,issued,published,container-title,publisher,type,score",
        }
    )
    request = Request(
        f"https://api.crossref.org/works?{parameters}",
        headers={"User-Agent": USER_AGENT, "Accept": "application/json"},
    )
    for attempt in range(3):
        try:
            with urlopen(request, timeout=25) as response:
                return json.load(response)["message"]["items"]
        except (HTTPError, URLError, TimeoutError):
            if attempt == 2:
                return []
            time.sleep(1.5 * (attempt + 1))
    return []


def candidate_year(candidate: dict) -> int | None:
    for key in ("published", "issued"):
        parts = candidate.get(key, {}).get("date-parts", [])
        if parts and parts[0] and parts[0][0] is not None:
            return int(parts[0][0])
    return None


def score_candidate(record: dict, candidate: dict) -> dict:
    candidate_title = clean_text((candidate.get("title") or [""])[0])
    title_similarity = SequenceMatcher(None, normalized(record["title"]), normalized(candidate_title)).ratio()
    surname = creator_surname(record["creator"])
    candidate_surnames = {normalized(author.get("family", "")) for author in candidate.get("author", [])}
    author_match = bool(surname and surname in candidate_surnames)
    source_year = parse_year(record["date"])
    match_year = candidate_year(candidate)
    year_match = source_year is not None and match_year is not None and abs(source_year - match_year) <= 1
    score = title_similarity * 0.84 + float(author_match) * 0.10 + float(year_match) * 0.06
    accepted = title_similarity >= 0.94 and (
        author_match or title_similarity >= 0.985
    ) and (
        source_year is None or match_year is None or abs(source_year - match_year) <= 2
    )
    container = candidate.get("container-title") or []
    return {
        "candidateTitle": candidate_title,
        "titleSimilarity": round(title_similarity, 4),
        "authorMatch": author_match,
        "sourceYear": source_year,
        "candidateYear": match_year,
        "yearMatch": year_match,
        "score": round(score, 4),
        "accepted": accepted,
        "doi": candidate.get("DOI", ""),
        "journal": clean_text(container[0]) if container else "",
        "publisher": clean_text(candidate.get("publisher", "")),
        "crossrefType": clean_text(candidate.get("type", "")),
        "authorCount": len(candidate.get("author", [])),
    }


def audit_record(record: dict) -> dict:
    scored = [score_candidate(record, candidate) for candidate in crossref_items(record)]
    best = max(scored, key=lambda candidate: candidate["score"], default=None)
    return {
        "id": record["id"],
        "title": record["title"],
        "creator": record["creator"],
        "date": record["date"],
        "sourceType": record["type"],
        "period": period(record["date"]),
        "match": best,
    }


def summarise(results: list[dict]) -> dict:
    accepted = [row for row in results if row["match"] and row["match"]["accepted"]]
    by_type: dict[str, dict] = {}
    for source_type in sorted({row["sourceType"] for row in results}):
        rows = [row for row in results if row["sourceType"] == source_type]
        matches = [row for row in rows if row["match"] and row["match"]["accepted"]]
        by_type[source_type] = {"sample": len(rows), "accepted": len(matches), "rate": len(matches) / len(rows)}
    return {
        "sample": len(results),
        "accepted": len(accepted),
        "acceptedRate": len(accepted) / len(results) if results else 0,
        "exactTitleMatches": sum(row["match"]["titleSimilarity"] == 1 for row in accepted),
        "journalAmongAccepted": sum(bool(row["match"]["journal"]) for row in accepted),
        "publisherAmongAccepted": sum(bool(row["match"]["publisher"]) for row in accepted),
        "multiAuthorAmongAccepted": sum(row["match"]["authorCount"] > 1 for row in accepted),
        "bySourceType": by_type,
        "acceptedCrossrefTypes": dict(Counter(row["match"]["crossrefType"] for row in accepted)),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "analysis" / "robustness-study" / "work" / "source",
    )
    parser.add_argument("--per-stratum", type=int, default=6)
    parser.add_argument("--seed", type=int, default=20260913)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "analysis" / "metadata-enrichment-pilot.json",
    )
    args = parser.parse_args()
    records = load_records(args.source.resolve())
    sample = draw_sample(records, args.per_stratum, args.seed)
    results = []
    with ThreadPoolExecutor(max_workers=6) as executor:
        futures = {executor.submit(audit_record, record): record for record in sample}
        for index, future in enumerate(as_completed(futures), start=1):
            results.append(future.result())
            print(f"[{index:03d}/{len(sample):03d}] {futures[future]['id']}", flush=True)
    results.sort(key=lambda row: row["id"])
    output = {
        "method": {
            "source": "Crossref REST API",
            "sampleDesign": f"up to {args.per_stratum} records per publication-type × period stratum",
            "seed": args.seed,
            "acceptance": "title similarity ≥ .94; author surname or near-exact title; year within two years when both are present",
        },
        "summary": summarise(results),
        "results": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(output["summary"], indent=2))


if __name__ == "__main__":
    main()
