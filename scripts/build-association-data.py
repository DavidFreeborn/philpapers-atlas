#!/usr/bin/env python3
"""Build metadata lenses and cross-lens association summaries.

The metadata workbook is aligned to the fixed 69,400-paper display cohort using
the same English-index and document-ID files as the map build. Missing metadata
is encoded as -1 and excluded pairwise from association calculations. HDBSCAN
noise remains a substantive category and is retained.
"""

from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PAPER_COUNT = 69_400
CATALOG_PATH = ROOT / "public" / "data" / "lenses" / "catalog.json"
LABELS_DIR = CATALOG_PATH.parent / "labels"
METADATA_IDS = {"metadata_publication_type", "metadata_publication_period"}

TYPE_CATEGORIES = [
    (0, "Article", "#56b4e9"),
    (1, "Book", "#e69f00"),
    (2, "Review", "#cc79a7"),
]
PERIOD_CATEGORIES = [
    (0, "Before 1990", "#8c72cb"),
    (1, "1990–1999", "#2aaac2"),
    (2, "2000–2009", "#2aaa4a"),
    (3, "2010–2019", "#e4d34f"),
    (4, "2020 onward", "#ea6a8a"),
    (5, "Manuscript or forthcoming", "#b262c2"),
]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_type(value: object) -> int:
    text = clean_text(value).casefold()
    if "review" in text:
        return 2
    if "book" in text or "книга" in text:
        return 1
    if "article" in text or "articolo" in text:
        return 0
    return -1


def normalize_period(value: object) -> int:
    text = clean_text(value).casefold()
    if not text:
        return -1
    if "manuscript" in text or "forthcoming" in text:
        return 5
    match = re.search(r"(?<!\d)(1[5-9]\d{2}|20\d{2})(?!\d)", text)
    if not match:
        return -1
    year = int(match.group(1))
    if year < 1990:
        return 0
    if year < 2000:
        return 1
    if year < 2010:
        return 2
    if year < 2020:
        return 3
    if year <= 2099:
        return 4
    return -1


def aligned_paper_ids(source: Path) -> list[str]:
    english_indices = np.load(source / "english_indices.npy", allow_pickle=False)
    document_paths = [
        line.strip()
        for line in (source / "document_ids.txt").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    if english_indices.shape != (PAPER_COUNT,):
        raise ValueError(f"Expected {PAPER_COUNT:,} English indices")
    if int(np.max(english_indices)) >= len(document_paths):
        raise ValueError("English index points past the document ID list")
    paper_ids = [Path(document_paths[int(index)]).stem for index in english_indices]
    if len(set(paper_ids)) != PAPER_COUNT:
        raise ValueError("Aligned paper identifiers are not unique")
    return paper_ids


def load_aligned_metadata(source: Path, paper_ids: list[str]) -> tuple[np.ndarray, np.ndarray]:
    needed_ids = set(paper_ids)
    workbook = load_workbook(source / "philpapers_metadata.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    columns = {name: index for index, name in enumerate(headers)}
    required = {"_id", "type", "date", "description"}
    if missing := required - columns.keys():
        raise ValueError(f"Metadata workbook is missing columns: {sorted(missing)}")

    records: dict[str, tuple[int, int, int]] = {}
    for values in rows:
        paper_id = clean_text(values[columns["_id"]])
        if not paper_id or paper_id not in needed_ids:
            continue
        candidate = (
            normalize_type(values[columns["type"]]),
            normalize_period(values[columns["date"]]),
            len(clean_text(values[columns["description"]])),
        )
        previous = records.get(paper_id)
        if previous is None or candidate[2] > previous[2]:
            records[paper_id] = candidate
    workbook.close()

    type_labels = np.full(PAPER_COUNT, -1, dtype="<i2")
    period_labels = np.full(PAPER_COUNT, -1, dtype="<i2")
    for index, paper_id in enumerate(paper_ids):
        record = records.get(paper_id)
        if record:
            type_labels[index] = record[0]
            period_labels[index] = record[1]
    return type_labels, period_labels


def category_rows(labels: np.ndarray, categories: list[tuple[int, str, str]]) -> list[dict]:
    return [
        {
            "id": category_id,
            "label": label,
            "count": int(np.count_nonzero(labels == category_id)),
            "terms": [],
            "color": colour,
        }
        for category_id, label, colour in categories
        if np.any(labels == category_id)
    ]


def metadata_lens(
    lens_id: str,
    name: str,
    option_label: str,
    field: str,
    labels: np.ndarray,
    categories: list[tuple[int, str, str]],
) -> dict:
    missing = int(np.count_nonzero(labels < 0))
    clusters = category_rows(labels, categories)
    return {
        "id": lens_id,
        "name": name,
        "optionLabel": option_label,
        "algorithm": "metadata",
        "preferred": False,
        "labelsFile": f"data/lenses/labels/{lens_id}.bin",
        "clusterCount": len(clusters),
        "noiseCount": missing,
        "noisePct": missing / PAPER_COUNT * 100,
        "method": {
            "metadataField": field,
            "sourceDescription": "PhilPapers metadata export",
        },
        "metrics": {
            "coveragePct": (PAPER_COUNT - missing) / PAPER_COUNT * 100,
        },
        "clusters": clusters,
    }


def include_label(lens: dict, labels: np.ndarray) -> np.ndarray:
    if lens["algorithm"] == "metadata":
        return labels >= 0
    return np.ones(labels.shape, dtype=bool)


def corrected_cramers_v(left: np.ndarray, right: np.ndarray) -> tuple[float, int]:
    if left.size <= 1:
        return 0.0, int(left.size)
    left_ids, left_inverse = np.unique(left, return_inverse=True)
    right_ids, right_inverse = np.unique(right, return_inverse=True)
    rows = left_ids.size
    columns = right_ids.size
    if rows <= 1 or columns <= 1:
        return 0.0, int(left.size)
    table = np.bincount(
        left_inverse * columns + right_inverse,
        minlength=rows * columns,
    ).reshape(rows, columns)
    total = int(table.sum())
    expected = table.sum(axis=1)[:, None] * table.sum(axis=0)[None, :] / total
    chi_square = float(np.sum(np.square(table - expected) / expected))
    phi_squared = chi_square / total
    correction = ((columns - 1) * (rows - 1)) / (total - 1)
    corrected_phi = max(0.0, phi_squared - correction)
    corrected_rows = rows - ((rows - 1) ** 2) / (total - 1)
    corrected_columns = columns - ((columns - 1) ** 2) / (total - 1)
    denominator = min(corrected_rows - 1, corrected_columns - 1)
    value = (corrected_phi / denominator) ** 0.5 if denominator > 0 else 0.0
    return value, total


def short_label(lens: dict) -> str:
    method = lens.get("method", {})
    if lens["algorithm"] == "hdbscan":
        dimensions = "2D" if method.get("clusteringSpace") == "display2d" else f"{method.get('umapDimensions')}D"
        return f"H{lens['clusterCount']}·{dimensions}"
    if lens["algorithm"] == "kmeans":
        return f"K{lens['clusterCount']}"
    if lens["algorithm"] == "lda":
        return f"LDA{lens['clusterCount']}"
    return "Type" if lens["id"] == "metadata_publication_type" else "Period"


def build_overview(catalog: dict, labels_by_id: dict[str, np.ndarray]) -> dict:
    lenses = catalog["lenses"]
    pairs = []
    for left_index, left_lens in enumerate(lenses):
        left_all = labels_by_id[left_lens["id"]]
        for right_lens in lenses[left_index + 1 :]:
            right_all = labels_by_id[right_lens["id"]]
            eligible = include_label(left_lens, left_all) & include_label(right_lens, right_all)
            value, count = corrected_cramers_v(left_all[eligible], right_all[eligible])
            pairs.append(
                {
                    "left": left_lens["id"],
                    "right": right_lens["id"],
                    "cramersV": round(value, 6),
                    "eligiblePapers": count,
                }
            )
    return {
        "version": "1.0.0",
        "paperCount": PAPER_COUNT,
        "measure": "bias-corrected Cramers V",
        "inclusion": "HDBSCAN noise is retained as a category; papers with unavailable metadata are excluded pairwise.",
        "lenses": [
            {"id": lens["id"], "label": lens["name"], "shortLabel": short_label(lens)}
            for lens in lenses
        ],
        "pairs": pairs,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "analysis" / "robustness-study" / "work" / "source",
    )
    args = parser.parse_args()
    source = args.source.resolve()
    for name in ("english_indices.npy", "document_ids.txt", "philpapers_metadata.xlsx"):
        if not (source / name).exists():
            raise FileNotFoundError(source / name)

    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    catalog["lenses"] = [lens for lens in catalog["lenses"] if lens["id"] not in METADATA_IDS]
    paper_ids = aligned_paper_ids(source)
    type_labels, period_labels = load_aligned_metadata(source, paper_ids)
    LABELS_DIR.mkdir(parents=True, exist_ok=True)
    type_labels.tofile(LABELS_DIR / "metadata_publication_type.bin")
    period_labels.tofile(LABELS_DIR / "metadata_publication_period.bin")

    catalog["lenses"].extend(
        [
            metadata_lens(
                "metadata_publication_type",
                "Publication type",
                "Publication type · 3 categories",
                "type",
                type_labels,
                TYPE_CATEGORIES,
            ),
            metadata_lens(
                "metadata_publication_period",
                "Publication period",
                "Publication period · 6 categories",
                "date",
                period_labels,
                PERIOD_CATEGORIES,
            ),
        ]
    )
    catalog["version"] = "3.0.0"
    CATALOG_PATH.write_text(
        json.dumps(catalog, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )

    labels_by_id = {
        lens["id"]: np.fromfile(ROOT / "public" / lens["labelsFile"], dtype="<i2")
        for lens in catalog["lenses"]
    }
    if any(labels.shape != (PAPER_COUNT,) for labels in labels_by_id.values()):
        raise ValueError("At least one lens label file is not aligned to the fixed cohort")
    overview = build_overview(catalog, labels_by_id)
    association_path = CATALOG_PATH.parent / "associations.json"
    association_path.write_text(
        json.dumps(overview, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        f"Built 2 metadata lenses and {len(overview['pairs'])} pairwise associations "
        f"for {PAPER_COUNT:,} papers."
    )


if __name__ == "__main__":
    main()
