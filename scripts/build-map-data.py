#!/usr/bin/env python3
"""Build compact, browser-ready PhilPapers Atlas data files.

The source artefacts are the saved preferred clustering outputs from Explorer.
No embeddings, PCA, UMAP, or HDBSCAN models are refit by this script.
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import math
import re
import shutil
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


DETAIL_CHUNK_SIZE = 500


def clean_text(value: object, limit: int | None = None) -> str:
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    if limit and len(text) > limit:
        text = text[: limit - 1].rstrip() + "…"
    return text


def clean_date(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = clean_text(value, 80)
    if re.fullmatch(r"\d{4}\.0", text):
        return text[:4]
    return text


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )


def load_cluster_summary(path: Path) -> list[dict[str, object]]:
    clusters: list[dict[str, object]] = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            clusters.append(
                {
                    "id": int(row["Cluster ID"]),
                    "label": clean_text(row["Label"]),
                    "count": int(row["Papers"]),
                    "terms": [
                        term.strip()
                        for term in row["TF-IDF Terms"].split(",")
                        if term.strip()
                    ],
                }
            )
    return sorted(clusters, key=lambda cluster: int(cluster["id"]))


def load_metadata(path: Path, needed_ids: set[str]) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(rows)]
    columns = {name: index for index, name in enumerate(headers)}
    required = {"_id", "title", "creator", "description", "date", "url"}
    missing = required - columns.keys()
    if missing:
        raise ValueError(f"Metadata workbook is missing columns: {sorted(missing)}")

    records: dict[str, dict[str, str]] = {}
    for values in rows:
        paper_id = clean_text(values[columns["_id"]])
        if not paper_id or paper_id not in needed_ids:
            continue
        candidate = {
            "id": paper_id,
            "title": clean_text(values[columns["title"]], 700),
            "authors": clean_text(values[columns["creator"]], 500),
            "abstract": clean_text(values[columns["description"]], 5000),
            "date": clean_date(values[columns["date"]]),
            "url": clean_text(values[columns["url"]], 500)
            or f"https://philpapers.org/rec/{paper_id}",
            "subject": clean_text(values[columns.get("subject", -1)], 300)
            if "subject" in columns
            else "",
        }
        previous = records.get(paper_id)
        if previous is None or len(candidate["abstract"]) > len(previous["abstract"]):
            records[paper_id] = candidate

    workbook.close()
    return records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    source = args.source.resolve()
    output = args.output.resolve()
    details_dir = output / "details"

    required_files = [
        "reduced.npy",
        "labels.npy",
        "english_indices.npy",
        "document_ids.txt",
        "philpapers_metadata.xlsx",
        "cluster_summary.csv",
        "cache_metadata.json",
    ]
    missing = [name for name in required_files if not (source / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing source files: {missing}")

    coordinates = np.load(source / "reduced.npy")
    labels = np.load(source / "labels.npy")
    english_indices = np.load(source / "english_indices.npy")
    document_paths = [
        line.strip()
        for line in (source / "document_ids.txt").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]

    if coordinates.shape != (len(labels), 2):
        raise ValueError(
            f"Coordinate shape {coordinates.shape} does not match {len(labels)} labels"
        )
    if len(english_indices) != len(labels):
        raise ValueError("English-index and label lengths differ")
    if int(np.max(english_indices)) >= len(document_paths):
        raise ValueError("English index points past the document ID list")

    paper_ids = [Path(document_paths[int(index)]).stem for index in english_indices]
    if len(set(paper_ids)) != len(paper_ids):
        raise ValueError("Paper identifiers are not unique after English filtering")

    metadata = load_metadata(source / "philpapers_metadata.xlsx", set(paper_ids))
    clusters = load_cluster_summary(source / "cluster_summary.csv")
    cluster_by_id = {int(cluster["id"]): cluster for cluster in clusters}

    output.mkdir(parents=True, exist_ok=True)
    if details_dir.exists():
        shutil.rmtree(details_dir)
    details_dir.mkdir(parents=True)

    search_records: list[list[object]] = []
    points: list[list[object]] = []
    detail_chunk: list[list[str]] = []
    missing_metadata = 0

    for record_index, (paper_id, coordinate, cluster_id) in enumerate(
        zip(paper_ids, coordinates, labels, strict=True)
    ):
        record = metadata.get(paper_id)
        if record is None:
            missing_metadata += 1
            record = {
                "id": paper_id,
                "title": paper_id,
                "authors": "",
                "abstract": "",
                "date": "",
                "url": f"https://philpapers.org/rec/{paper_id}",
                "subject": "",
            }

        title = record["title"] or paper_id
        search_records.append(
            [paper_id, title, record["authors"], record["date"]]
        )
        points.append(
            [
                round(float(coordinate[0]), 5),
                round(float(coordinate[1]), 5),
                int(cluster_id),
                record_index,
            ]
        )
        detail_chunk.append(
            [
                paper_id,
                title,
                record["authors"],
                record["date"],
                record["url"],
                record["abstract"],
                record["subject"],
            ]
        )
        if len(detail_chunk) == DETAIL_CHUNK_SIZE or record_index == len(labels) - 1:
            chunk_index = record_index // DETAIL_CHUNK_SIZE
            write_json(details_dir / f"{chunk_index:03d}.json", detail_chunk)
            detail_chunk = []

    actual_counts: dict[int, int] = {}
    for label in labels.tolist():
        actual_counts[int(label)] = actual_counts.get(int(label), 0) + 1
    for cluster in clusters:
        cluster_id = int(cluster["id"])
        if cluster_id not in actual_counts:
            raise ValueError(f"Cluster summary contains absent cluster {cluster_id}")
        cluster["count"] = actual_counts[cluster_id]

    unknown_cluster_ids = sorted(
        set(actual_counts) - {-1} - set(cluster_by_id)
    )
    if unknown_cluster_ids:
        raise ValueError(f"Missing cluster descriptions: {unknown_cluster_ids}")

    finite = np.isfinite(coordinates)
    if not bool(finite.all()):
        raise ValueError("The saved 2D projection contains non-finite coordinates")

    bounds = {
        "minX": float(np.min(coordinates[:, 0])),
        "maxX": float(np.max(coordinates[:, 0])),
        "minY": float(np.min(coordinates[:, 1])),
        "maxY": float(np.max(coordinates[:, 1])),
    }
    pipeline = json.loads((source / "cache_metadata.json").read_text(encoding="utf-8"))

    write_json(
        output / "map.json",
        {
            "version": "candidate-42",
            "count": len(points),
            "bounds": bounds,
            "points": points,
        },
    )
    write_json(output / "search.json", search_records)
    write_json(
        output / "clusters.json",
        {
            "noiseCount": actual_counts.get(-1, 0),
            "clusters": clusters,
        },
    )
    write_json(
        output / "manifest.json",
        {
            "version": "candidate-42",
            "papers": len(points),
            "clusteredPapers": len(points) - actual_counts.get(-1, 0),
            "noisePapers": actual_counts.get(-1, 0),
            "clusterCount": len(clusters),
            "detailChunkSize": DETAIL_CHUNK_SIZE,
            "detailChunks": math.ceil(len(points) / DETAIL_CHUNK_SIZE),
            "metadataMatched": len(points) - missing_metadata,
            "metadataMissing": missing_metadata,
            "pipeline": pipeline,
        },
    )

    print(
        json.dumps(
            {
                "papers": len(points),
                "clusters": len(clusters),
                "noise": actual_counts.get(-1, 0),
                "metadata_missing": missing_metadata,
                "output": str(output),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
